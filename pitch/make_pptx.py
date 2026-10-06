"""Generate the ACRES pitch deck as an editable .pptx file.

Run:  .venv/bin/python pitch/make_pptx.py
Output: pitch/acres_pitch_deck.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x18, 0x41, 0x8A)
NAVY_DARK = RGBColor(0x10, 0x2E, 0x63)
ORANGE = RGBColor(0xF2, 0x65, 0x22)
TEAL = RGBColor(0x00, 0x9E, 0x73)
INK = RGBColor(0x2A, 0x34, 0x42)
MUTED = RGBColor(0x5A, 0x64, 0x72)
GRAY = RGBColor(0xF4, 0xF6, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Helvetica Neue"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_left_bar(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), SLIDE_H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    bar.shadow.inherit = False


def add_text(slide, left, top, width, height, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """runs: list of (text, size_pt, bold, color, space_after_pt)."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for idx, (text, size, bold, color, space_after) in enumerate(runs):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = text
        p.alignment = align
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.space_after = Pt(space_after)
    return box


def add_kicker(slide, text):
    add_text(slide, Inches(0.9), Inches(0.7), Inches(11.5), Inches(0.5),
             [(text.upper(), 14, True, ORANGE, 0)])


def add_title(slide, text):
    add_text(slide, Inches(0.9), Inches(1.15), Inches(11.6), Inches(1.4),
             [(text, 34, True, NAVY, 0)])


def add_bullets(slide, bullets, top=Inches(2.9), size=18):
    runs = [(f"•   {b}", size, False, INK, 10) for b in bullets]
    add_text(slide, Inches(0.95), top, Inches(11.4), Inches(4.2), runs)


def add_card(slide, left, top, width, height, heading, body, accent):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = GRAY
    card.line.fill.background()
    card.shadow.inherit = False
    card.adjustments[0] = 0.06

    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top + Inches(0.08),
                                   Inches(0.09), height - Inches(0.16))
    strip.fill.solid()
    strip.fill.fore_color.rgb = accent
    strip.line.fill.background()
    strip.shadow.inherit = False

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.15)
    p = tf.paragraphs[0]
    p.text = heading
    p.font.name = FONT
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(3)
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.name = FONT
    p2.font.size = Pt(13)
    p2.font.color.rgb = INK


def add_cards(slide, cards, top=Inches(2.9)):
    """cards: list of (heading, body, accent). Rendered in two columns."""
    left_x = Inches(0.95)
    col_w = Inches(5.5)
    gap = Inches(0.5)
    row_h = Inches(1.5)
    row_gap = Inches(0.35)
    for idx, (heading, body, accent) in enumerate(cards):
        col = idx % 2
        row = idx // 2
        x = left_x + col * (col_w + gap)
        y = top + row * (row_h + row_gap)
        add_card(slide, x, y, col_w, row_h, heading, body, accent)


def add_pills(slide, labels, top=Inches(3.0)):
    x = Inches(0.95)
    for label in labels:
        w = Inches(0.32 * len(label) + 0.6)
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, w, Inches(0.55))
        pill.fill.solid()
        pill.fill.fore_color.rgb = GRAY
        pill.line.color.rgb = RGBColor(0xDD, 0xE3, 0xEA)
        pill.shadow.inherit = False
        tf = pill.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.text = label
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = NAVY
        x += w + Inches(0.3)


def add_footer(slide, n, total):
    add_text(slide, Inches(11.6), Inches(6.95), Inches(1.3), Inches(0.4),
             [(f"{n} / {total}", 10, False, MUTED, 0)], align=PP_ALIGN.RIGHT)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    slides = []

    # 1. Title
    s = prs.slides.add_slide(blank)
    add_left_bar(s)
    add_text(s, Inches(0.9), Inches(0.9), Inches(6), Inches(0.5),
             [("ACRES", 16, True, ORANGE, 0)])
    add_text(s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(2.2),
             [("ACRES Clearance Automated Intake", 46, True, NAVY, 0)])
    add_text(s, Inches(0.9), Inches(3.9), Inches(10.8), Inches(1.6),
             [("A self-serve portal that turns a messy clearance request into a "
               "ready-to-submit ACRES row and routes it straight to the project "
               "executive — no middleman.", 20, False, INK, 0)])
    add_pills(s, ["Self-serve", "AI-assisted", "Working prototype"], top=Inches(5.6))
    slides.append(s)

    # 2. Problem
    s = prs.slides.add_slide(blank); add_left_bar(s)
    add_kicker(s, "The problem")
    add_title(s, "Clearance is slow, fragmented and error-prone")
    add_bullets(s, [
        "A complex, 20+ column Excel sheet must be filled perfectly.",
        "Project details — ops manager, department, clearance level — are scattered.",
        "Applicants manually decode clearance categories and MSD reference numbers.",
        "Requests bounce through a middleman before the project executive.",
        "Wrong or incomplete rows cause rework and delay.",
        "Sensitive data is handled ad hoc over email.",
    ])
    slides.append(s)

    # 3. Cost
    s = prs.slides.add_slide(blank); add_left_bar(s)
    add_kicker(s, "The cost")
    add_title(s, "Every hop and typo adds days")
    add_cards(s, [
        ("Days of waiting", "Per clearance, multiplied across a project.", ORANGE),
        ("Rework loops", "Between applicant, middleman and executive.", ORANGE),
        ("Risk", "Sensitive details scattered across email and spreadsheets.", ORANGE),
        ("No standard", "Every request is formatted differently.", ORANGE),
    ])
    slides.append(s)

    # 4. Solution
    s = prs.slides.add_slide(blank); add_left_bar(s)
    add_kicker(s, "The solution")
    add_title(s, "A self-serve portal that does the work for you")
    add_bullets(s, [
        "Paste an email, or upload a PDF / image clearance document.",
        "AI extracts and validates every required field.",
        "Project details auto-fill from a project register.",
        "Produces a standardized, ready-to-submit ACRES row.",
        "Packages it and routes it directly to the project executive.",
    ])
    slides.append(s)

    # 5. How it works
    s = prs.slides.add_slide(blank); add_left_bar(s)
    add_kicker(s, "How it works")
    add_title(s, "From messy input to routed request")
    add_cards(s, [
        ("1 · Input", "Paste text or upload a scanned PDF/image.", NAVY),
        ("2 · Read", "OCR turns documents into text.", NAVY),
        ("3 · Understand", "A language model extracts fields, constrained to a strict schema.", NAVY),
        ("4 · Verify", "Deterministic parsing for MSD M/…, categories like 1A/02/2B, and dates.", TEAL),
        ("5 · Assemble", "Merge with the project register into one ACRES row.", TEAL),
        ("6 · Route", "Package and send to the project executive.", ORANGE),
    ])
    slides.append(s)

    # 6. Where AI is used
    s = prs.slides.add_slide(blank); add_left_bar(s)
    add_kicker(s, "Where AI is used — and where it isn't")
    add_title(s, "AI for language. Code for critical values.")
    add_cards(s, [
        ("AI does", "Reading documents, extracting fields, drafting the “Nature of Involvement” sentence.", NAVY),
        ("Deterministic code does", "MSD reference numbers, clearance categories, dates, project lookup, validation, the final row.", TEAL),
        ("Human-in-the-loop", "Anything missing or mis-formatted is flagged for the user to fix.", ORANGE),
    ])
    slides.append(s)

    # 7. Key features
    s = prs.slides.add_slide(blank); add_left_bar(s)
    add_kicker(s, "Key features")
    add_title(s, "Everything needed, built in")
    add_bullets(s, [
        "AI Intake Assistant — paste anything, the form fills itself.",
        "Document OCR — scanned PDFs and photos of forms.",
        "Smart conditionals — locals vs foreigners; visit-reference country rules.",
        "Format validation — NRIC, dates, mobile, email.",
        "Project auto-fill — location, end date, ops manager, clearance level, department, executive.",
        "Direct routing — recipient pre-filled; one-click email package.",
        "Flexible AI backend — local (private) or cloud; free tiers supported.",
    ], size=16)
    slides.append(s)

    # 8. Differentiators
    s = prs.slides.add_slide(blank); add_left_bar(s)
    add_kicker(s, "Why it's different")
    add_title(s, "It removes the middleman — not just digitises the form")
    add_cards(s, [
        ("Hybrid intelligence", "LLM for language; deterministic code for critical values.", NAVY),
        ("Privacy options", "Can run fully local — clearance data never leaves the machine.", TEAL),
        ("Self-serve", "Applicants complete it themselves, correctly, first time.", ORANGE),
        ("Standardized", "Same validated ACRES row for every project and person type.", NAVY),
    ])
    slides.append(s)

    # 9. Who it's for
    s = prs.slides.add_slide(blank); add_left_bar(s)
    add_kicker(s, "Who it's for")
    add_title(s, "Built for everyone in the clearance chain")
    add_pills(s, ["Applicants", "Project executives", "Ops managers", "Project / security admin"], top=Inches(3.1))
    add_text(s, Inches(0.95), Inches(4.3), Inches(11.4), Inches(1.4),
             [("Applicants complete the form; executives receive a ready package; "
               "ops managers endorse; admin teams get consistency and an audit trail.",
               16, False, INK, 0)])
    slides.append(s)

    # 10. Impact
    s = prs.slides.add_slide(blank); add_left_bar(s)
    add_kicker(s, "Impact")
    add_title(s, "Faster, cleaner, safer")
    add_cards(s, [
        ("Faster", "Minutes instead of days of back-and-forth.", TEAL),
        ("More accurate", "Validated, standardized rows; fewer rework loops.", TEAL),
        ("Less admin", "The middleman step is eliminated.", NAVY),
        ("Safer & scalable", "Privacy-preserving local AI; one flow for every project.", NAVY),
    ])
    slides.append(s)

    # 11. Status
    s = prs.slides.add_slide(blank); add_left_bar(s)
    add_kicker(s, "Status")
    add_title(s, "A working prototype, not a slide")
    add_bullets(s, [
        "Deployed and iterated with real feedback.",
        "Reads real scanned clearance documents (OCR + AI).",
        "Handles the full ACRES field set and per-project routing.",
        "AI backend options integrated: local Ollama, Gemini free tier, OpenCode Go, any OpenAI-compatible provider.",
    ], size=17)
    slides.append(s)

    # 12. Roadmap
    s = prs.slides.add_slide(blank); add_left_bar(s)
    add_kicker(s, "Roadmap")
    add_title(s, "From prototype to production")
    add_bullets(s, [
        "Real email sending (SMTP/API) with delivery receipts.",
        "Live project-register integration and SSO / role-based access.",
        "Audit trail and submission tracking.",
        "Multiple clearance rows per applicant; category-aware routing.",
        "Analytics: time-to-clearance, error rates, bottlenecks.",
    ])
    slides.append(s)

    # 13. Risks
    s = prs.slides.add_slide(blank); add_left_bar(s)
    add_kicker(s, "Risks & mitigations")
    add_title(s, "Handled deliberately")
    add_cards(s, [
        ("Sensitive data", "Local model option; access codes; minimal retention.", ORANGE),
        ("Model accuracy", "Schema constraints + deterministic parsing + human review.", ORANGE),
        ("Adoption", "Works alongside the existing Excel template; low training.", NAVY),
        ("Process variation", "Routing is configurable per project.", NAVY),
    ])
    slides.append(s)

    # 14. Ask
    s = prs.slides.add_slide(blank); add_left_bar(s)
    add_kicker(s, "The ask")
    add_text(s, Inches(0.9), Inches(1.2), Inches(11.5), Inches(1.6),
             [("Let's clear the middleman.", 40, True, NAVY, 0)])
    add_bullets(s, [
        "Pilot with one project team.",
        "Gather feedback from applicants, ops managers and project executives.",
        "Measure time-to-clearance and rework, before and after.",
        "Expand to more projects.",
    ], top=Inches(3.1))
    slides.append(s)

    total = len(slides)
    for n, s in enumerate(slides, start=1):
        add_footer(s, n, total)

    out = Path(__file__).with_name("acres_pitch_deck.pptx")
    prs.save(out)
    print(f"wrote {out} ({total} slides)")


if __name__ == "__main__":
    build()
